"""Build YC-style pitch deck for Friday. Run: python docs/presentation/build_pitch_deck.py"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent / "Friday-Voice-Agent.pptx"

BLACK = RGBColor(0x0A, 0x0A, 0x0A)
WHITE = RGBColor(0xFA, 0xFA, 0xFA)
GRAY = RGBColor(0xA1, 0xA1, 0xAA)
ACCENT = RGBColor(0xF9, 0x73, 0x16)  # warm orange


def _bg(slide, color=BLACK):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _title(slide, text, y=2.0, size=44, color=WHITE, bold=True):
    box = slide.shapes.add_textbox(Inches(0.75), Inches(y), Inches(11.5), Inches(1.4))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = PP_ALIGN.LEFT


def _subtitle(slide, text, y=3.4, size=22):
    box = slide.shapes.add_textbox(Inches(0.75), Inches(y), Inches(11.5), Inches(1.2))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = GRAY


def _bullets(slide, title, items, title_y=0.55):
    _title(slide, title, y=title_y, size=32, color=ACCENT)
    box = slide.shapes.add_textbox(Inches(0.85), Inches(1.45), Inches(11.2), Inches(5.2))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(20)
        p.font.color.rgb = WHITE
        p.space_after = Pt(14)


def _diagram_flow(slide):
    _title(slide, "How Friday works", y=0.45, size=30, color=ACCENT)
    steps = [
        ("You", "Voice or text"),
        ("Friday", "Route + tools"),
        ("Map + GitHub", "Ground truth"),
        ("You", "Yes to post"),
    ]
    x0 = 0.6
    w = 2.6
    gap = 0.35
    y = 2.35
    h = 1.35
    for i, (head, sub) in enumerate(steps):
        x = x0 + i * (w + gap)
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0x1C, 0x1C, 0x1E)
        shape.line.color.rgb = ACCENT
        tf = shape.text_frame
        tf.paragraphs[0].text = head
        tf.paragraphs[0].font.size = Pt(18)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = WHITE
        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(13)
        p2.font.color.rgb = GRAY
        if i < len(steps) - 1:
            ax = x + w + 0.05
            arrow = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, Inches(ax), Inches(y + 0.45), Inches(gap - 0.1), Inches(0.35)
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = GRAY
            arrow.line.fill.background()
    _subtitle(
        slide,
        "Jev picks map tools · Gateway loop for multi-step GitHub · Writes never auto-post",
        y=4.15,
        size=16,
    )


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "Friday", y=2.3, size=52)
    _subtitle(s, "Stay in control when code—and threats—move faster than review.", y=3.35, size=24)

    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "We are in an existential squeeze.", y=2.0, size=38)
    _subtitle(
        s,
        "AI generates code at machine speed.\nHumans are still asked to review all of it.\nThat math does not work.",
        y=3.2,
        size=22,
    )

    s = prs.slides.add_slide(blank)
    _bg(s)
    _bullets(
        s,
        "The review bottleneck",
        [
            "PR queues and diff mountains grow with every agent-assisted sprint",
            "LGTM becomes survival, not assurance",
            "Nobody holds architecture + history + issues in one head",
        ],
    )

    s = prs.slides.add_slide(blank)
    _bg(s)
    _bullets(
        s,
        "The security asymmetry",
        [
            "AI finds vulnerabilities and unsafe patterns quickly",
            "Exploitation automates; understanding does not",
            "Ship fast without orientation → lose grip on what actually runs",
        ],
    )

    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "Friday", y=1.8, size=44, color=ACCENT)
    _subtitle(
        s,
        "A voice-first control surface for the repo you have open.\n"
        "See the shape. Ask out loud. Stage GitHub actions. Confirm before you post.",
        y=2.9,
        size=22,
    )

    s = prs.slides.add_slide(blank)
    _bg(s)
    _bullets(
        s,
        "Product",
        [
            "Circle-pack map — entry points, core, hotspots",
            "Landmarks + layers — heat, imports, focus",
            "Voice (AssemblyAI) + typed Ask — same brain",
            "Commits, issues, PRs, security scan, merge review",
        ],
    )

    s = prs.slides.add_slide(blank)
    _bg(s)
    _diagram_flow(s)

    s = prs.slides.add_slide(blank)
    _bg(s)
    _bullets(
        s,
        "Human in the loop",
        [
            "create_issue / add_comment → draft only",
            "Separate turn: yes → confirm_write posts",
            "No silent auto-merge, no drive-by comments",
            "OAuth + tokens stay server-side",
        ],
    )

    s = prs.slides.add_slide(blank)
    _bg(s)
    _bullets(
        s,
        "Judgment, not vibes",
        [
            "Jev: confident tool routing for map questions",
            "Review gate: structured safe-to-merge narrative",
            "Escalate when confidence is low (Vertex / K2)",
            "Honest give-up when the repo cannot answer",
        ],
    )

    s = prs.slides.add_slide(blank)
    _bg(s)
    _bullets(
        s,
        "Stack",
        [
            "FastAPI · AssemblyAI Voice + STT + LLM Gateway",
            "GitHub REST / OAuth · D3 map ingest",
            "Deployed on Google Cloud Run",
        ],
    )

    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "Try it", y=2.0, size=40, color=ACCENT)
    _subtitle(
        s,
        "https://friday-voice-agent-147606977567.us-central1.run.app\n"
        "?repo=owner/name\n\n"
        "Hold the mic · Explain the map · Check the latest commit",
        y=3.0,
        size=20,
    )

    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "Why now", y=1.9, size=36)
    _subtitle(
        s,
        "The window where humans still approve what ships is narrowing.\n"
        "Friday is how you keep orientation and consent in that window.",
        y=3.0,
        size=22,
    )

    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "Friday", y=2.5, size=48)
    _subtitle(s, "github.com/hatif03/friday-voice-agent", y=3.4, size=22)

    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
